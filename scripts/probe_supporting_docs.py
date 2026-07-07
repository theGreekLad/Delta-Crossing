#!/usr/bin/env python3
"""Probe Supporting Docs files and report extractable text by type."""

from __future__ import annotations

import json
import re
import zipfile
import xml.etree.ElementTree as ET
from collections import defaultdict
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
DOCS = ROOT / "Supporting Docs"


def clean(text: str, limit: int = 400) -> str:
    text = re.sub(r"\s+", " ", text or "").strip()
    return text[:limit] + ("..." if len(text) > limit else "")


def read_pdf(path: Path) -> str:
    import fitz

    doc = fitz.open(path)
    chunks = []
    for page in doc:
        chunks.append(page.get_text("text"))
        if sum(len(part) for part in chunks) > 2000:
            break
    doc.close()
    return "\n".join(chunks)


def read_docx(path: Path) -> str:
    from docx import Document

    doc = Document(path)
    return "\n".join(p.text for p in doc.paragraphs if p.text.strip())


def read_xlsx(path: Path) -> str:
    from openpyxl import load_workbook

    wb = load_workbook(path, read_only=True, data_only=True)
    lines = []
    for sheet in wb.worksheets[:3]:
        lines.append(f"[Sheet: {sheet.title}]")
        for row_idx, row in enumerate(sheet.iter_rows(values_only=True)):
            values = [str(cell).strip() for cell in row if cell not in (None, "")]
            if values:
                lines.append(" | ".join(values))
            if row_idx >= 12:
                break
    wb.close()
    return "\n".join(lines)


def read_pptx(path: Path) -> str:
    from pptx import Presentation

    prs = Presentation(path)
    lines = []
    for slide_idx, slide in enumerate(prs.slides[:8], start=1):
        slide_text = []
        for shape in slide.shapes:
            if hasattr(shape, "text") and shape.text.strip():
                slide_text.append(shape.text.strip())
        if slide_text:
            lines.append(f"[Slide {slide_idx}] " + " / ".join(slide_text))
    return "\n".join(lines)


def read_ods(path: Path) -> str:
    from odf.opendocument import load
    from odf.table import Table, TableRow, TableCell
    from odf.text import P

    doc = load(str(path))
    lines = []
    for table in doc.spreadsheet.getElementsByType(Table)[:2]:
        lines.append(f"[Table: {table.getAttribute('name')}]")
        for row_idx, row in enumerate(table.getElementsByType(TableRow)):
            cells = []
            for cell in row.getElementsByType(TableCell):
                parts = []
                for paragraph in cell.getElementsByType(P):
                    parts.append(str(paragraph))
                cells.append(" ".join(parts).strip())
            if any(cells):
                lines.append(" | ".join(cells))
            if row_idx >= 12:
                break
    return "\n".join(lines)


def read_dwg(path: Path) -> str:
    raise NotImplementedError("DWG is a binary CAD format; no parser installed")


def read_heic(path: Path) -> str:
    try:
        from pillow_heif import register_heif_opener
        from PIL import Image

        register_heif_opener()
        with Image.open(path) as image:
            return f"Image readable: {image.size[0]}x{image.size[1]} {image.mode} (no OCR run)"
    except ImportError as exc:
        raise NotImplementedError("HEIC requires pillow-heif for image open; OCR not installed") from exc


READERS = {
    ".pdf": read_pdf,
    ".docx": read_docx,
    ".xlsx": read_xlsx,
    ".pptx": read_pptx,
    ".ods": read_ods,
    ".dwg": read_dwg,
    ".heic": read_heic,
}


def main() -> None:
    if not DOCS.exists():
        print(json.dumps({"error": f"Missing folder: {DOCS}"}, indent=2))
        return

    by_ext: dict[str, list[Path]] = defaultdict(list)
    for path in sorted(DOCS.rglob("*")):
        if path.is_file():
            by_ext[path.suffix.lower()].append(path)

    summary = {
        "folder": str(DOCS),
        "total_files": sum(len(paths) for paths in by_ext.values()),
        "by_extension": {},
    }

    for ext, paths in sorted(by_ext.items()):
        reader = READERS.get(ext)
        sample = paths[0]
        entry = {
            "count": len(paths),
            "readable": reader is not None,
            "sample_file": str(sample.relative_to(ROOT)).replace("\\", "/"),
        }
        if reader:
            try:
                text = reader(sample)
                entry["sample_chars"] = len(text)
                entry["sample_preview"] = clean(text, 500)
                entry["status"] = "ok" if text.strip() else "empty"
            except NotImplementedError as exc:
                entry["status"] = "unsupported"
                entry["note"] = str(exc)
            except Exception as exc:  # noqa: BLE001
                entry["status"] = "error"
                entry["note"] = str(exc)
        else:
            entry["status"] = "no_reader"
            entry["note"] = "No extractor configured for this extension"
        summary["by_extension"][ext or "(none)"] = entry

    print(json.dumps(summary, indent=2))


if __name__ == "__main__":
    main()
