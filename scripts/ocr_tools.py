#!/usr/bin/env python3
"""OCR helpers for scanned PDFs and images in Supporting Docs."""

from __future__ import annotations

import hashlib
import json
import re
from dataclasses import dataclass
from datetime import datetime, timezone
from pathlib import Path

import fitz
import numpy as np

ROOT = Path(__file__).resolve().parents[1]
DOCS_ROOT = ROOT / "Supporting Docs"
CACHE_ROOT = DOCS_ROOT / "_ocr_cache"
HASH_INDEX_FILE = CACHE_ROOT / "hash_index.json"
OCR_EXTENSIONS = {".pdf", ".heic", ".jpg", ".jpeg", ".png", ".tif", ".tiff"}
TEXT_EXTENSIONS = {".docx", ".xlsx", ".ods", ".pptx"}
SKIP_NAME_PARTS = (
    "market study",
    "delta apartments market study",
)


def should_skip_file(path: Path) -> bool:
    name = path.name.lower()
    return any(part in name for part in SKIP_NAME_PARTS)

_reader = None


def get_reader():
    global _reader
    if _reader is None:
        import easyocr

        _reader = easyocr.Reader(["en"], gpu=False, verbose=False)
    return _reader


def cache_path_for(source: Path) -> Path:
    rel = source.relative_to(DOCS_ROOT)
    return CACHE_ROOT / rel.with_suffix(rel.suffix + ".ocr.json")


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def load_hash_index() -> dict[str, str]:
    if not HASH_INDEX_FILE.exists():
        return {}
    try:
        return json.loads(HASH_INDEX_FILE.read_text(encoding="utf-8"))
    except (OSError, json.JSONDecodeError):
        return {}


def save_hash_index(index: dict[str, str]) -> None:
    HASH_INDEX_FILE.parent.mkdir(parents=True, exist_ok=True)
    HASH_INDEX_FILE.write_text(json.dumps(index, indent=2), encoding="utf-8")


def reuse_hash_cache(source: Path, digest: str) -> DocumentText | None:
    index = load_hash_index()
    cached_rel = index.get(digest)
    if not cached_rel:
        return None
    cached_file = ROOT / cached_rel
    if not cached_file.exists():
        return None
    payload = json.loads(cached_file.read_text(encoding="utf-8"))
    payload["source"] = str(source.relative_to(ROOT)).replace("\\", "/")
    payload["signature"] = file_signature(source)
    payload["content_hash"] = digest
    cache_file = save_cache(source, payload)
    return DocumentText(
        source=source,
        method=payload.get("method", "cached_hash"),
        text=payload.get("full_text", ""),
        cached=True,
        cache_file=cache_file,
    )


def register_hash_cache(source: Path, digest: str, cache_file: Path) -> None:
    index = load_hash_index()
    index[digest] = str(cache_file.relative_to(ROOT)).replace("\\", "/")
    save_hash_index(index)


def choose_max_pages(source: Path, requested: int | None, batch_cap: int = 50) -> int | None:
    if requested is not None:
        return requested
    if source.suffix.lower() != ".pdf":
        return None

    doc = fitz.open(source)
    page_count = doc.page_count
    doc.close()

    native_preview = extract_native_text(source, max_pages=min(3, page_count))
    if not needs_ocr(source, native_preview):
        return None

    if page_count <= batch_cap:
        return None
    return batch_cap


def file_signature(path: Path) -> dict:
    stat = path.stat()
    return {
        "size": stat.st_size,
        "mtime": int(stat.st_mtime),
    }


def load_cache(source: Path) -> dict | None:
    cache_file = cache_path_for(source)
    if not cache_file.exists():
        return None
    try:
        payload = json.loads(cache_file.read_text(encoding="utf-8"))
    except (OSError, json.JSONDecodeError):
        return None
    if payload.get("signature") != file_signature(source):
        return None
    return payload


def save_cache(source: Path, payload: dict) -> Path:
    cache_file = cache_path_for(source)
    cache_file.parent.mkdir(parents=True, exist_ok=True)
    cache_file.write_text(json.dumps(payload, indent=2), encoding="utf-8")
    return cache_file


def pixmap_to_array(pix: fitz.Pixmap) -> np.ndarray:
    channels = pix.n
    array = np.frombuffer(pix.samples, dtype=np.uint8).reshape(pix.height, pix.width, channels)
    if channels == 4:
        return array[:, :, :3]
    return array


def ocr_image_array(image: np.ndarray) -> str:
    reader = get_reader()
    lines = reader.readtext(image)
    return "\n".join(text.strip() for _, text, _ in lines if text and text.strip())


def ocr_pdf(path: Path, max_pages: int | None = None, scale: float = 2.0) -> dict:
    doc = fitz.open(path)
    pages = []
    full_parts = []
    page_count = doc.page_count if max_pages is None else min(doc.page_count, max_pages)

    for index in range(page_count):
        page = doc[index]
        pix = page.get_pixmap(matrix=fitz.Matrix(scale, scale), alpha=False)
        text = ocr_image_array(pixmap_to_array(pix))
        pages.append({"page": index + 1, "text": text})
        if text.strip():
            full_parts.append(text)

    doc.close()
    return {
        "pages": pages,
        "full_text": "\n\n".join(full_parts),
        "page_count": page_count,
    }


def ocr_image_file(path: Path) -> dict:
    if path.suffix.lower() == ".heic":
        try:
            from pillow_heif import register_heif_opener
            from PIL import Image

            register_heif_opener()
            image = np.array(Image.open(path).convert("RGB"))
        except ImportError as exc:
            raise RuntimeError("HEIC OCR requires pillow-heif: pip install pillow-heif") from exc
    else:
        from PIL import Image

        image = np.array(Image.open(path).convert("RGB"))

    text = ocr_image_array(image)
    return {
        "pages": [{"page": 1, "text": text}],
        "full_text": text,
        "page_count": 1,
    }


def extract_native_text(path: Path, max_pages: int | None = None) -> str:
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        doc = fitz.open(path)
        limit = doc.page_count if max_pages is None else min(doc.page_count, max_pages)
        chunks = [doc[index].get_text("text") for index in range(limit)]
        doc.close()
        return "\n".join(chunks)
    if suffix == ".docx":
        from docx import Document

        return "\n".join(p.text for p in Document(path).paragraphs if p.text.strip())
    if suffix == ".xlsx":
        from openpyxl import load_workbook

        wb = load_workbook(path, read_only=True, data_only=True)
        lines = []
        for sheet in wb.worksheets:
            lines.append(f"[Sheet: {sheet.title}]")
            for row in sheet.iter_rows(values_only=True):
                values = [str(cell).strip() for cell in row if cell not in (None, "")]
                if values:
                    lines.append(" | ".join(values))
        wb.close()
        return "\n".join(lines)
    if suffix == ".ods":
        from odf.opendocument import load
        from odf.table import Table, TableCell, TableRow
        from odf.text import P

        doc = load(str(path))
        lines = []
        for table in doc.spreadsheet.getElementsByType(Table):
            lines.append(f"[Table: {table.getAttribute('name')}]")
            for row in table.getElementsByType(TableRow):
                cells = []
                for cell in row.getElementsByType(TableCell):
                    cells.append(" ".join(str(p) for p in cell.getElementsByType(P)).strip())
                if any(cells):
                    lines.append(" | ".join(cells))
        return "\n".join(lines)
    if suffix == ".pptx":
        import zipfile
        import xml.etree.ElementTree as ET

        lines = []
        with zipfile.ZipFile(path) as archive:
            slide_names = sorted(
                name for name in archive.namelist() if name.startswith("ppt/slides/slide") and name.endswith(".xml")
            )
            ns = {"a": "http://schemas.openxmlformats.org/drawingml/2006/main"}
            for slide_name in slide_names:
                root = ET.fromstring(archive.read(slide_name))
                texts = [node.text for node in root.findall(".//a:t", ns) if node.text]
                if texts:
                    lines.append("[Slide] " + " / ".join(texts))
        return "\n".join(lines)
    return ""


def needs_ocr(path: Path, native_text: str) -> bool:
    if path.suffix.lower() in OCR_EXTENSIONS:
        return len(re.sub(r"\s+", "", native_text)) < 80
    return False


@dataclass
class DocumentText:
    source: Path
    method: str
    text: str
    cached: bool
    cache_file: Path | None = None


def extract_document_text(
    source: Path,
    *,
    force_ocr: bool = False,
    max_pages: int | None = None,
    use_cache: bool = True,
) -> DocumentText:
    if not source.exists():
        raise FileNotFoundError(source)

    digest = sha256_file(source)
    if use_cache and not force_ocr:
        reused = reuse_hash_cache(source, digest)
        if reused:
            return reused

    if source.suffix.lower() in TEXT_EXTENSIONS:
        return DocumentText(
            source=source,
            method="native",
            text=extract_native_text(source, max_pages=max_pages),
            cached=False,
        )

    if source.suffix.lower() not in OCR_EXTENSIONS:
        raise ValueError(f"Unsupported file type: {source.suffix}")

    preview = extract_native_text(source, max_pages=3)
    if use_cache and not force_ocr:
        cached = load_cache(source)
        if cached:
            return DocumentText(
                source=source,
                method=cached.get("method", "ocr"),
                text=cached.get("full_text", ""),
                cached=True,
                cache_file=cache_path_for(source),
            )

    if not force_ocr and not needs_ocr(source, preview):
        native_text = extract_native_text(source, max_pages=None)
        payload = {
            "source": str(source.relative_to(ROOT)).replace("\\", "/"),
            "signature": file_signature(source),
            "content_hash": digest,
            "method": "native_pdf_text",
            "engine": "pymupdf",
            "generated_at": datetime.now(timezone.utc).isoformat(),
            "full_text": native_text,
            "pages": [{"page": 1, "text": native_text}],
            "page_count": 1,
        }
        cache_file = save_cache(source, payload)
        register_hash_cache(source, digest, cache_file)
        return DocumentText(source=source, method="native_pdf_text", text=native_text, cached=False, cache_file=cache_file)

    if source.suffix.lower() == ".pdf":
        ocr_payload = ocr_pdf(source, max_pages=max_pages)
        method = "ocr_pdf"
    else:
        ocr_payload = ocr_image_file(source)
        method = "ocr_image"

    payload = {
        "source": str(source.relative_to(ROOT)).replace("\\", "/"),
        "signature": file_signature(source),
        "content_hash": digest,
        "method": method,
        "engine": "easyocr",
        "generated_at": datetime.now(timezone.utc).isoformat(),
        **ocr_payload,
    }
    if max_pages is not None and ocr_payload.get("page_count", 0) >= max_pages:
        payload["truncated"] = True
        payload["max_pages"] = max_pages
    cache_file = save_cache(source, payload)
    register_hash_cache(source, digest, cache_file)
    return DocumentText(
        source=source,
        method=method,
        text=ocr_payload["full_text"],
        cached=False,
        cache_file=cache_file,
    )


def iter_supporting_docs(root: Path = DOCS_ROOT) -> list[Path]:
    files = []
    for path in sorted(root.rglob("*")):
        if not path.is_file():
            continue
        if path.relative_to(root).parts[0].startswith("_"):
            continue
        if should_skip_file(path):
            continue
        if path.suffix.lower() in OCR_EXTENSIONS | TEXT_EXTENSIONS:
            files.append(path)
    return files


def read_cached_text(source: Path) -> DocumentText | None:
    cached = load_cache(source)
    if cached:
        return DocumentText(
            source=source,
            method=cached.get("method", "cached"),
            text=cached.get("full_text", ""),
            cached=True,
            cache_file=cache_path_for(source),
        )

    if source.suffix.lower() in TEXT_EXTENSIONS:
        text = extract_native_text(source)
        return DocumentText(source=source, method="native", text=text, cached=False)

    if source.suffix.lower() == ".pdf":
        text = extract_native_text(source)
        if len(re.sub(r"\s+", "", text)) >= 80:
            return DocumentText(source=source, method="native_pdf_text", text=text, cached=False)

    return None


def search_cached_docs(query: str, root: Path = DOCS_ROOT, limit: int = 20) -> list[dict]:
    pattern = re.compile(re.escape(query), re.IGNORECASE)
    hits = []
    for source in iter_supporting_docs(root):
        try:
            doc = read_cached_text(source)
            if doc is None or not doc.text:
                continue
        except Exception:
            continue
        for match in pattern.finditer(doc.text):
            start = max(0, match.start() - 80)
            end = min(len(doc.text), match.end() + 120)
            hits.append(
                {
                    "file": str(source.relative_to(ROOT)).replace("\\", "/"),
                    "method": doc.method,
                    "snippet": doc.text[start:end].replace("\n", " "),
                }
            )
            if len(hits) >= limit:
                return hits
    return hits
